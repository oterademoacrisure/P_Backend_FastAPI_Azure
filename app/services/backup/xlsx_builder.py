"""
Assembles the final .xlsx deliverable for the LangGraph pipeline
(app/graph.py's finalize_node) from the model's generated plain-text draft.

The generated text follows the convention documented in
app/services/prompt_templates.py: one or more '## <Section Title>' headed
sections, each either plain prose/bullets or '|'-delimited rows (header row
first). This builder turns that into one worksheet per section -- a '|'-row
section becomes a table (bold header row), a prose section becomes wrapped
text in a single column.
"""

from __future__ import annotations

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

OUTPUT_DIR = os.getenv("GENERATED_OUTPUT_DIR", "generated_outputs")

_SECTION_HEADING_RE = re.compile(r"^##\s*(.+?)\s*$", re.MULTILINE)


def split_sections(content: str) -> list[tuple[str, str]]:
    """Splits the draft into (title, body) pairs on '## ' headings. A draft
    with no headings at all becomes a single ("Content", draft) section.
    Public so app/services/draft_repair.py can parse drafts the same way."""
    matches = list(_SECTION_HEADING_RE.finditer(content))
    if not matches:
        return [("Content", content.strip())]

    sections = []
    for i, m in enumerate(matches):
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
        sections.append((m.group(1), content[start:end].strip()))
    return sections


def _safe_sheet_title(title: str, used: set[str]) -> str:
    # Excel sheet names: <=31 chars, no []:*?/\\
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _write_table_section(ws: Worksheet, body: str) -> None:
    rows = [line for line in body.splitlines() if line.strip()]
    for r, line in enumerate(rows, start=1):
        cells = [c.strip() for c in line.split("|")]
        for c, value in enumerate(cells, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if r == 1:
                cell.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 40


def _write_prose_section(ws: Worksheet, body: str) -> None:
    lines = body.splitlines() or [""]
    for r, line in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def is_tabular(body: str) -> bool:
    data_lines = [line for line in body.splitlines() if line.strip()]
    if not data_lines:
        return False
    return sum(1 for line in data_lines if "|" in line) / len(data_lines) >= 0.5


def build_output(output_format: str, content: str, session_id: str) -> str:
    """Writes `content` to a .xlsx file under OUTPUT_DIR and returns its path."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()

    for title, body in split_sections(content):
        ws = wb.create_sheet(title=_safe_sheet_title(title, used_titles))
        if is_tabular(body):
            _write_table_section(ws, body)
        else:
            _write_prose_section(ws, body)

    filename = f"{session_id}_{output_format}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return path
