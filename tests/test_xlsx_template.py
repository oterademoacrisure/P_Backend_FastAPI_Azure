"""
Unit tests for xlsx_builder.fill_template() against the real common
templates/STTM_Data_Ingestion_Template.xlsx -- no network.
"""

from __future__ import annotations

import os
import tempfile
from datetime import date

import openpyxl

from app import xlsx_builder

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "templates", "STTM_Data_Ingestion_Template.xlsx")

DRAFT = """## 1. STTM Summary
Attribute|Detail
Report / Pipeline Name|Vendor3 Medical Claims
Version / Date|1.0 / 2024-06-01
Extra Attribute|Not in the template

## 2. STTM Mapping
Mapping ID|Target Table / Output Object (Entity Name)|Target Field Name|Source Field|Mapping Confidence
M-001|Claim Header|Claim_ID|Claim Number|Confirmed
M-002|Claim Line|Units|Units|Confirmed
M-003|Claim Line|Modifier_1|Modifier 1|Candidate

## 3. Assumptions and Open Q
Type|ID|Statement|Basis / Note
Open Question|Q-001|Confirm Paid Date on denied lines|Vendor dictionary vs sample

## 4. SME Review Checklist
Review Area|Review Question|Status
Privacy|Are PHI fields handled?|Pending SME Review
"""


def _fill(draft: str = DRAFT):
    with open(TEMPLATE, "rb") as f:
        template = f.read()
    path = os.path.join(tempfile.mkdtemp(), "out.xlsx")
    xlsx_builder.fill_template(template, draft, path)
    return openpyxl.load_workbook(path)


class TestFillSttmTemplate:
    def test_keeps_template_sheets_and_banner(self):
        wb = _fill()
        assert wb.sheetnames == ["1. STTM Summary", "2. STTM Mapping", "3. Assumptions and Open Q", "4. SME Review Checklist"]
        summary = wb["1. STTM Summary"]
        assert summary["A1"].value.startswith("COGNIZANT")
        assert "A1:B1" in [str(r) for r in summary.merged_cells.ranges]

    def test_summary_values_go_next_to_template_labels(self):
        summary = _fill()["1. STTM Summary"]
        values = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(3, summary.max_row + 1)}
        assert values["Report / Pipeline Name"] == "Vendor3 Medical Claims"
        assert values["Version / Date"] == f"1.0 / {date.today().isoformat()}"
        assert values["Extra Attribute"] == "Not in the template"

    def test_mapping_values_land_in_matching_template_columns(self):
        mapping = _fill()["2. STTM Mapping"]
        header = [c.value for c in mapping[1]]
        assert len([h for h in header if h]) == 22  # the template's own columns, unchanged
        row = {header[c - 1]: mapping.cell(2, c).value for c in range(1, len(header) + 1)}
        assert row["Mapping ID"] == "M-001"
        assert row["Target Table / Output Object (Entity Name)"] == "Claim Header"
        assert row["Source Field"] == "Claim Number"
        assert row["Mapping Confidence"] == "Confirmed"
        assert row["Report Section"] is None  # not generated, left blank rather than shifted

    def test_banding_continues_past_preformatted_rows(self):
        lines = "\n".join(f"M-{i:03d}|Claim Line|Units|Units|Confirmed" for i in range(1, 31))
        draft = "## 2. STTM Mapping\nMapping ID|Target Table / Output Object (Entity Name)|Target Field Name|Source Field|Mapping Confidence\n" + lines
        mapping = _fill(draft)["2. STTM Mapping"]
        assert mapping.cell(31, 1).value == "M-030"
        assert mapping.cell(2, 1).fill.fgColor.rgb == mapping.cell(30, 1).fill.fgColor.rgb  # even rows shaded
        assert mapping.cell(3, 1).fill.fgColor.rgb == mapping.cell(31, 1).fill.fgColor.rgb
        assert mapping.freeze_panes == "A2"

    def test_other_sheets_filled(self):
        wb = _fill()
        assert wb["3. Assumptions and Open Q"].cell(2, 3).value == "Confirm Paid Date on denied lines"
        assert wb["4. SME Review Checklist"].cell(2, 1).value == "Privacy"

    def test_unknown_section_gets_its_own_sheet(self):
        wb = _fill(DRAFT + "\n## 5. Lineage Notes\nSome prose the template has no sheet for.")
        assert "5. Lineage Notes" in wb.sheetnames
